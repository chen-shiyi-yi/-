"""多格式报告渲染器 - 支持 Markdown、HTML、PDF、DOCX"""

from __future__ import annotations

from pathlib import Path

from jinja2 import Environment, FileSystemLoader

from src.report.builder import ReportBuilder
from src.report.exporter import ReportExporter


class ReportRenderer:
    """报告渲染器"""

    def __init__(self, template_dir: Path | None = None):
        if template_dir is None:
            template_dir = Path(__file__).parent / "templates"
        self.template_dir = template_dir
        self.env = Environment(
            loader=FileSystemLoader(str(template_dir)),
            keep_trailing_newline=True,
        )
        self.exporter = ReportExporter()

    def render_markdown(self, builder: ReportBuilder, output_path: Path) -> Path:
        """渲染 Markdown 格式报告"""
        template = self.env.get_template("report.md.j2")
        context = builder.build_context()
        content = template.render(**context)
        output_file = output_path.with_suffix(".md")
        output_file.parent.mkdir(parents=True, exist_ok=True)
        output_file.write_text(content, encoding="utf-8")
        return output_file

    def render_html(self, builder: ReportBuilder, output_path: Path) -> Path:
        """渲染 HTML 格式报告"""
        template = self.env.get_template("report.html.j2")
        context = builder.build_context()
        content = template.render(**context)
        output_file = output_path.with_suffix(".html")
        output_file.parent.mkdir(parents=True, exist_ok=True)
        output_file.write_text(content, encoding="utf-8")
        return output_file

    def render_pdf(self, builder: ReportBuilder, output_path: Path) -> Path:
        """渲染 PDF 格式报告（通过 HTML 转换）"""
        from weasyprint import HTML

        html_content = self._render_html_string(builder)
        output_file = output_path.with_suffix(".pdf")
        output_file.parent.mkdir(parents=True, exist_ok=True)
        HTML(string=html_content).write_pdf(str(output_file))
        return output_file

    def render_docx(self, builder: ReportBuilder, output_path: Path) -> Path:
        """渲染 DOCX 格式报告"""
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        context = builder.build_context()
        doc = Document()

        # 标题
        title = doc.add_heading(context["project"].project_name, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 生成信息
        doc.add_paragraph(f"生成时间: {context['generated_at']}")
        doc.add_paragraph("")

        # 执行摘要
        doc.add_heading("执行摘要", level=1)
        doc.add_paragraph(context["executive_summary"])

        # 统计概览
        doc.add_heading("发现统计", level=1)
        stats = context["stats"]
        table = doc.add_table(rows=2, cols=5)
        table.style = "Table Grid"
        headers = ["严重", "高危", "中危", "低危", "信息"]
        values = [
            str(stats["critical"]),
            str(stats["high"]),
            str(stats["medium"]),
            str(stats["low"]),
            str(stats["info"]),
        ]
        for i, (h, v) in enumerate(zip(headers, values)):
            table.rows[0].cells[i].text = h
            table.rows[1].cells[i].text = v

        # 资产清单
        doc.add_heading("资产清单", level=1)
        for host in context["hosts"]:
            doc.add_heading(f"{host.ip} ({host.hostname})" if host.hostname else host.ip, level=2)
            if host.os:
                doc.add_paragraph(f"操作系统: {host.os}")
            if host.services:
                svc_table = doc.add_table(rows=1, cols=4)
                svc_table.style = "Table Grid"
                for i, h in enumerate(["端口", "协议", "服务", "版本"]):
                    svc_table.rows[0].cells[i].text = h
                for svc in host.services:
                    row = svc_table.add_row()
                    row.cells[0].text = str(svc.port)
                    row.cells[1].text = svc.protocol
                    row.cells[2].text = svc.service
                    row.cells[3].text = svc.version

        # 漏洞详情
        doc.add_heading("漏洞详情", level=1)
        for i, vuln in enumerate(context["vulnerabilities"], 1):
            severity_label = context["severity_labels"].get(vuln.severity.value, "")
            doc.add_heading(f"{i}. [{severity_label}] {vuln.name}", level=2)

            if vuln.cvss_score:
                doc.add_paragraph(f"CVSS 分数: {vuln.cvss_score}")
            if vuln.cve_ids:
                doc.add_paragraph(f"CVE: {', '.join(vuln.cve_ids)}")
            if vuln.owasp_category:
                doc.add_paragraph(f"OWASP: {vuln.owasp_category}")
            doc.add_paragraph(f"受影响资产: {vuln.affected_asset}")

            doc.add_heading("漏洞描述", level=3)
            doc.add_paragraph(vuln.description)

            if vuln.reproduction_steps:
                doc.add_heading("复现步骤", level=3)
                for j, step in enumerate(vuln.reproduction_steps, 1):
                    doc.add_paragraph(f"{j}. {step}")

            if vuln.impact:
                doc.add_heading("影响分析", level=3)
                doc.add_paragraph(vuln.impact)

            if vuln.remediation:
                doc.add_heading("修复建议", level=3)
                doc.add_paragraph(vuln.remediation)

            doc.add_paragraph("")

        output_file = output_path.with_suffix(".docx")
        output_file.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_file))
        return output_file

    def render_all(self, builder: ReportBuilder, output_path: Path) -> dict[str, Path]:
        """渲染所有格式"""
        results = {}
        results["md"] = self.render_markdown(builder, output_path)
        results["html"] = self.render_html(builder, output_path)
        try:
            results["pdf"] = self.render_pdf(builder, output_path)
        except Exception as e:
            results["pdf_error"] = str(e)
        results["docx"] = self.render_docx(builder, output_path)
        return results

    def export_json(self, builder: ReportBuilder, output_path: Path) -> Path:
        """导出为 JSON 格式"""
        return self.exporter.export_json(builder.report, output_path)

    def export_csv(self, builder: ReportBuilder, output_path: Path) -> Path:
        """导出为 CSV 格式"""
        return self.exporter.export_csv(builder.report, output_path)

    def export_summary(self, builder: ReportBuilder, output_path: Path) -> Path:
        """导出摘要报告"""
        return self.exporter.export_summary(builder.report, output_path)

    def _render_html_string(self, builder: ReportBuilder) -> str:
        template = self.env.get_template("report.html.j2")
        context = builder.build_context()
        return template.render(**context)
